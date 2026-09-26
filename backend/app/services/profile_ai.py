"""Phase 5b — AI Profile Understanding (plan v5; TDD agent A03).

resume file → text → one AI call → schema-validated structured profile, stored in `evidence_claims`
tagged as **self-reported** evidence (lowest weight — context for scoring, never a score on its own).
The candidate then confirms or corrects it; the confirmed version is what later phases use.
"""
import io
import logging
from typing import List, Optional

from docx import Document
from pydantic import BaseModel, Field, field_validator
from pypdf import PdfReader

from app import database
from app.ai import provider
from app.content.registry import get_type
from app.models.common import new_id, utcnow
from app.services import audit, files
from app.services.content import Conflict, NotFound, get_current

log = logging.getLogger(__name__)

PROMPT_KEY = "profile_parser"
MAX_CHARS = 24000
MIN_CHARS = 150


def _clean_list(values: List[str], limit: int) -> List[str]:
    seen, out = set(), []
    for v in values:
        v = " ".join(str(v).split())[:80]
        if v and v.lower() not in seen:
            seen.add(v.lower())
            out.append(v)
    return out[:limit]


class Role(BaseModel):
    title: str = Field(min_length=1, max_length=120)
    organisation: Optional[str] = Field(None, max_length=120)
    start: Optional[str] = Field(None, max_length=20)
    end: Optional[str] = Field(None, max_length=20)
    highlights: List[str] = Field(default_factory=list, max_length=6)

    @field_validator("highlights")
    @classmethod
    def _h(cls, v):
        return [" ".join(str(x).split())[:240] for x in v if str(x).strip()][:6]


class Education(BaseModel):
    qualification: str = Field(min_length=1, max_length=120)
    institution: Optional[str] = Field(None, max_length=120)
    year: Optional[int] = Field(None, ge=1950, le=2100)


class ResumeProfile(BaseModel):
    summary: Optional[str] = Field(None, max_length=600)
    total_years_experience: Optional[float] = Field(None, ge=0, le=60)
    roles: List[Role] = Field(default_factory=list, max_length=20)
    industries: List[str] = Field(default_factory=list)
    skills: List[str] = Field(default_factory=list)
    education: List[Education] = Field(default_factory=list, max_length=10)

    @field_validator("industries")
    @classmethod
    def _ind(cls, v):
        return _clean_list(v, 15)

    @field_validator("skills")
    @classmethod
    def _sk(cls, v):
        return _clean_list(v, 40)


def extract_text(doc: dict) -> str:
    data = files.read_bytes(doc)
    if doc["kind"] == "pdf":
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages[:15])
    elif doc["kind"] == "docx":
        d = Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:
        raise ValueError("unsupported resume type")
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())[:MAX_CHARS]


def _claim_out(c: dict) -> dict:
    return {
        "id": c["_id"],
        "status": c["status"],  # processing | parsed | needs_review | unreadable | unavailable | confirmed
        "message": c.get("message"),
        "evidence_tier": c["evidence_tier"],
        "file": c.get("file"),
        "parsed": c.get("parsed"),
        "confirmed": c.get("confirmed"),
        "model": c.get("model"),
        "prompt_version": c.get("prompt_version"),
        "created_at": c["created_at"],
        "confirmed_at": c.get("confirmed_at"),
    }


async def start(tenant_id: str, user_id: str, file_doc: dict) -> dict:
    """Record the resume on the account and open a claim in 'processing'; call `run` in the background."""
    await database.users().update_one({"_id": user_id, "tenant_id": tenant_id}, {"$set": {"resume_file_id": file_doc["_id"], "updated_at": utcnow()}})
    claim = {
        "_id": new_id(), "tenant_id": tenant_id, "user_id": user_id, "source": "resume", "evidence_tier": "self_reported",
        "file_id": file_doc["_id"], "file": files.public(file_doc), "status": "processing", "parsed": None, "confirmed": None,
        "created_at": utcnow(), "updated_at": utcnow(),
    }
    await database.evidence_claims().insert_one(claim)
    await audit.record(tenant_id=tenant_id, actor_id=user_id, candidate_id=user_id, entity_type="resume", entity_id=file_doc["_id"], action="resume.uploaded")
    return _claim_out(claim)


async def run(tenant_id: str, claim_id: str) -> None:
    """Parse the resume for one claim. Never raises — the outcome is written to the claim."""
    claim = await database.evidence_claims().find_one({"_id": claim_id, "tenant_id": tenant_id})
    if claim is None:
        return
    update = {"updated_at": utcnow()}
    try:
        text = extract_text(await files.get(tenant_id, claim["file_id"]))
        if len(text) < MIN_CHARS:
            update.update(status="unreadable", message="We couldn't read text from this file (it may be a scanned image). Please fill in your details below, or upload a text-based PDF or DOCX.")
        else:
            prompt = await get_current(get_type("prompts"), tenant_id, PROMPT_KEY)
            policy = prompt.data.get("model_policy") or {}
            parsed = await provider.generate(
                prompt.data["system_prompt"],
                f"Resume text:\n<<<\n{text}\n>>>",
                ResumeProfile,
                temperature=float(policy.get("temperature", 0.1)),
            )
            update.update(status="parsed", parsed=parsed.model_dump(), model=provider.model_label(), prompt_version=prompt.version, message=None)
    except NotFound:
        update.update(status="unavailable", message="Resume reading isn't set up yet — please fill in your details below.")
    except provider.AIUnavailable as exc:
        log.warning("resume parse unavailable", extra={"claim_id": claim_id, "reason": str(exc)})
        update.update(status="unavailable", message="We couldn't read your resume automatically right now — please fill in your details below.")
    except provider.AIOutputInvalid as exc:
        log.warning("resume parse invalid", extra={"claim_id": claim_id})
        update.update(status="needs_review", message="We couldn't reliably read your resume — please fill in or correct the details below.", raw_output=exc.raw)
    except Exception:  # noqa: BLE001 — a bad file must never leave the claim stuck in "processing"
        log.exception("resume parse failed", extra={"claim_id": claim_id})
        update.update(status="unreadable", message="We couldn't open this file. Please upload a different PDF or DOCX, or fill in your details below.")
    await database.evidence_claims().update_one({"_id": claim_id, "status": "processing"}, {"$set": update})
    await audit.record(
        tenant_id=tenant_id, actor_id="system", candidate_id=claim["user_id"], entity_type="evidence_claim", entity_id=claim_id,
        action="resume.parsed", metadata={"status": update["status"], "model": update.get("model")},
    )


async def latest(tenant_id: str, user_id: str) -> Optional[dict]:
    claim = await database.evidence_claims().find_one({"tenant_id": tenant_id, "user_id": user_id, "source": "resume"}, sort=[("created_at", -1)])
    return _claim_out(claim) if claim else None


async def confirm(tenant_id: str, user_id: str, claim_id: str, profile: ResumeProfile) -> dict:
    claim = await database.evidence_claims().find_one({"_id": claim_id, "tenant_id": tenant_id, "user_id": user_id})
    if claim is None:
        raise NotFound("Resume profile not found")
    if claim["status"] == "processing":
        raise Conflict("We're still reading your resume — try again in a moment")
    now = utcnow()
    edited = claim.get("parsed") is None or profile.model_dump() != claim.get("parsed")
    await database.evidence_claims().update_one(
        {"_id": claim_id},
        {"$set": {"status": "confirmed", "confirmed": profile.model_dump(), "confirmed_at": now, "edited_by_candidate": edited, "updated_at": now}},
    )
    await audit.record(
        tenant_id=tenant_id, actor_id=user_id, candidate_id=user_id, entity_type="evidence_claim", entity_id=claim_id,
        action="resume.confirmed", metadata={"edited": edited},
    )
    return _claim_out(await database.evidence_claims().find_one({"_id": claim_id}))


async def retry(tenant_id: str, user_id: str, claim_id: str) -> dict:
    claim = await database.evidence_claims().find_one({"_id": claim_id, "tenant_id": tenant_id, "user_id": user_id})
    if claim is None:
        raise NotFound("Resume profile not found")
    if claim["status"] not in ("unavailable", "needs_review", "unreadable"):
        raise Conflict("Nothing to retry")
    await database.evidence_claims().update_one({"_id": claim_id}, {"$set": {"status": "processing", "message": None, "updated_at": utcnow()}})
    return _claim_out(await database.evidence_claims().find_one({"_id": claim_id}))
